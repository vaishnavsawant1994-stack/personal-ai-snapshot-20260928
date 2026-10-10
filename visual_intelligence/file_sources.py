from __future__ import annotations

import base64
import csv
import io
import json
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader


_TEXT_EXTENSIONS = {
    '.txt', '.md', '.markdown', '.rst', '.csv', '.json', '.jsonl', '.yaml', '.yml',
    '.py', '.js', '.jsx', '.ts', '.tsx', '.html', '.css', '.scss', '.sql', '.toml',
    '.ini', '.cfg', '.xml', '.java', '.kt', '.swift', '.go', '.rs', '.c', '.h', '.cpp',
}


def decode_file_payload(filename: str, content_base64: str, *, max_bytes: int = 6 * 1024 * 1024) -> bytes:
    try:
        raw = base64.b64decode(str(content_base64 or ''), validate=True)
    except Exception as exc:
        raise ValueError('File payload is not valid base64') from exc
    if not raw:
        raise ValueError('File is empty')
    if len(raw) > max_bytes:
        raise ValueError(f'File exceeds the {max_bytes // (1024 * 1024)} MB Visualize limit')
    if not str(filename or '').strip():
        raise ValueError('File name is required')
    return raw


def _pdf_text(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    return '\n\n'.join((page.extract_text() or '').strip() for page in reader.pages if (page.extract_text() or '').strip())


def _docx_text(raw: bytes) -> str:
    document = Document(io.BytesIO(raw))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _csv_text(raw: bytes) -> str:
    decoded = raw.decode('utf-8-sig', errors='replace')
    rows = list(csv.reader(io.StringIO(decoded)))
    if not rows:
        return ''
    return '\n'.join(' | '.join(cell.strip() for cell in row) for row in rows[:1000])


def extract_source_text(filename: str, raw: bytes, *, max_chars: int = 20000) -> tuple[str, dict]:
    suffix = PurePath(filename).suffix.casefold()
    if suffix == '.pdf':
        text = _pdf_text(raw)
        kind = 'pdf'
    elif suffix == '.docx':
        text = _docx_text(raw)
        kind = 'docx'
    elif suffix == '.csv':
        text = _csv_text(raw)
        kind = 'csv'
    elif suffix in _TEXT_EXTENSIONS:
        text = raw.decode('utf-8-sig', errors='replace')
        kind = suffix.lstrip('.') or 'text'
    else:
        raise ValueError(f'Unsupported file type: {suffix or "unknown"}. Use PDF, DOCX, CSV, JSON, Markdown, text, or source code.')
    text = text.strip()
    if not text:
        raise ValueError('No readable text could be extracted from this file')
    return text[:max_chars], {
        'filename': filename,
        'kind': kind,
        'bytes': len(raw),
        'truncated': len(text) > max_chars,
        'extracted_chars': min(len(text), max_chars),
    }
