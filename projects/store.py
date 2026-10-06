from __future__ import annotations

import json
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path
from threading import RLock


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


class ProjectStore:
    """Small SQLite repository for the single-owner project workspace."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = RLock()
        self._init()

    def con(self):
        con = sqlite3.connect(self.path, timeout=10)
        con.row_factory = sqlite3.Row
        con.execute('PRAGMA foreign_keys=ON')
        return con

    def _init(self):
        with self.con() as con:
            con.executescript('''
              CREATE TABLE IF NOT EXISTS projects(
                id TEXT PRIMARY KEY, name TEXT NOT NULL, goal TEXT NOT NULL DEFAULT '',
                description TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'active',
                target_date TEXT, conversation_id TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_tasks(
                id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                title TEXT NOT NULL, description TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'planned',
                priority TEXT NOT NULL DEFAULT 'medium', owner TEXT NOT NULL DEFAULT 'owner', due_date TEXT,
                milestone_id TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_milestones(
                id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                title TEXT NOT NULL, outcome TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'planned',
                target_date TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_threads(
                id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                title TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'open', created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_replies(
                id TEXT PRIMARY KEY, thread_id TEXT NOT NULL REFERENCES project_threads(id) ON DELETE CASCADE,
                author TEXT NOT NULL, content TEXT NOT NULL, created_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_files(
                id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                title TEXT NOT NULL, kind TEXT NOT NULL DEFAULT 'link', url TEXT NOT NULL,
                knowledge_id TEXT, created_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_file_blobs(
                file_id TEXT PRIMARY KEY REFERENCES project_files(id) ON DELETE CASCADE,
                media_type TEXT NOT NULL, content BLOB NOT NULL, size_bytes INTEGER NOT NULL);
              CREATE TABLE IF NOT EXISTS project_file_versions(
                file_id TEXT NOT NULL REFERENCES project_files(id) ON DELETE CASCADE,
                version INTEGER NOT NULL, media_type TEXT NOT NULL, content BLOB NOT NULL,
                size_bytes INTEGER NOT NULL, created_at TEXT NOT NULL, created_by TEXT NOT NULL DEFAULT 'owner',
                PRIMARY KEY(file_id,version));
              CREATE TABLE IF NOT EXISTS project_file_index(
                file_id TEXT PRIMARY KEY REFERENCES project_files(id) ON DELETE CASCADE,
                project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                extracted_text TEXT NOT NULL DEFAULT '', state TEXT NOT NULL DEFAULT 'not_indexed',
                error TEXT NOT NULL DEFAULT '', indexed_at TEXT);
              CREATE TABLE IF NOT EXISTS project_activity(
                id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                actor TEXT NOT NULL, action TEXT NOT NULL, target_type TEXT NOT NULL, target_id TEXT,
                detail TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_task_files(
                task_id TEXT NOT NULL REFERENCES project_tasks(id) ON DELETE CASCADE,
                file_id TEXT NOT NULL REFERENCES project_files(id) ON DELETE CASCADE,
                project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                PRIMARY KEY(task_id,file_id));
              CREATE TABLE IF NOT EXISTS project_approvals(
                id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                title TEXT NOT NULL, summary TEXT NOT NULL, impact_summary TEXT NOT NULL DEFAULT '',
                scope_summary TEXT NOT NULL DEFAULT '', proposed_by TEXT NOT NULL DEFAULT 'owner',
                proposed_at TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending',
                reviewer TEXT, reviewed_at TEXT, decision TEXT, reviewer_comments TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS project_approval_tasks(
                approval_id TEXT NOT NULL REFERENCES project_approvals(id) ON DELETE CASCADE,
                task_id TEXT NOT NULL REFERENCES project_tasks(id) ON DELETE CASCADE,
                PRIMARY KEY(approval_id,task_id));
              CREATE TABLE IF NOT EXISTS project_approval_files(
                approval_id TEXT NOT NULL REFERENCES project_approvals(id) ON DELETE CASCADE,
                file_id TEXT NOT NULL REFERENCES project_files(id) ON DELETE CASCADE,
                PRIMARY KEY(approval_id,file_id));
              CREATE TABLE IF NOT EXISTS project_approval_history(
                id TEXT PRIMARY KEY, approval_id TEXT NOT NULL REFERENCES project_approvals(id) ON DELETE CASCADE,
                actor TEXT NOT NULL, decision TEXT NOT NULL, comments TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL);
              CREATE INDEX IF NOT EXISTS idx_projects_updated ON projects(updated_at DESC);
              CREATE INDEX IF NOT EXISTS idx_project_tasks_project ON project_tasks(project_id, created_at);
              CREATE INDEX IF NOT EXISTS idx_project_activity_project ON project_activity(project_id, created_at DESC);
            ''')
            # Additive migrations keep existing owner workspaces intact.
            additions = {
                'projects': {
                    'project_type': "TEXT NOT NULL DEFAULT 'software'",
                    'success_criteria': "TEXT NOT NULL DEFAULT ''",
                    'instructions': "TEXT NOT NULL DEFAULT ''",
                    'context_notes': "TEXT NOT NULL DEFAULT ''",
                    'is_walkthrough': "INTEGER NOT NULL DEFAULT 0",
                },
                'project_threads': {
                    'linked_item_type': "TEXT NOT NULL DEFAULT ''",
                    'linked_item_id': 'TEXT',
                    'pinned_reply_id': 'TEXT',
                    'followed': 'INTEGER NOT NULL DEFAULT 0',
                    'resolved_by': 'TEXT',
                    'resolved_at': 'TEXT',
                },
                'project_files': {
                    'folder_id': 'TEXT',
                    'is_pinned': 'INTEGER NOT NULL DEFAULT 0',
                    'source_state': "TEXT NOT NULL DEFAULT 'stored'",
                    'deliverable_status': "TEXT NOT NULL DEFAULT ''",
                    'source_provider': "TEXT NOT NULL DEFAULT ''",
                    'source_file_id': "TEXT NOT NULL DEFAULT ''",
                    'source_modified_at': "TEXT NOT NULL DEFAULT ''",
                },
                'project_milestones': {
                    'completion_criteria': "TEXT NOT NULL DEFAULT ''",
                },
                'project_tasks': {
                    'context_notes': "TEXT NOT NULL DEFAULT ''",
                    'execution_workflow_id': 'TEXT',
                    'execution_run_id': 'TEXT',
                    'execution_result': "TEXT NOT NULL DEFAULT ''",
                },
            }
            for table, columns in additions.items():
                present = {row['name'] for row in con.execute(f'PRAGMA table_info({table})')}
                for name, declaration in columns.items():
                    if name not in present:
                        con.execute(f'ALTER TABLE {table} ADD COLUMN {name} {declaration}')
            con.execute('''CREATE TABLE IF NOT EXISTS project_milestone_tasks(
                milestone_id TEXT NOT NULL REFERENCES project_milestones(id) ON DELETE CASCADE,
                task_id TEXT NOT NULL REFERENCES project_tasks(id) ON DELETE CASCADE,
                project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                PRIMARY KEY(milestone_id, task_id))''')
            con.execute('''INSERT OR IGNORE INTO project_milestone_tasks(milestone_id,task_id,project_id)
                SELECT milestone_id,id,project_id FROM project_tasks WHERE milestone_id IS NOT NULL''')
            try:
                con.execute("CREATE VIRTUAL TABLE IF NOT EXISTS project_file_fts USING fts5(file_id UNINDEXED, project_id UNINDEXED, title, body)")
                con.execute("CREATE INDEX IF NOT EXISTS idx_project_file_versions ON project_file_versions(file_id,version DESC)")
            except sqlite3.OperationalError:
                # Some packaged SQLite builds omit FTS5; the scoped LIKE fallback remains functional.
                pass
            for blob in con.execute('''SELECT f.id,f.project_id,f.title,b.media_type,b.content,b.size_bytes
                    FROM project_files f JOIN project_file_blobs b ON b.file_id=f.id
                    LEFT JOIN project_file_index i ON i.file_id=f.id WHERE i.file_id IS NULL''').fetchall():
                text = self._extract_text(blob['title'], blob['media_type'], bytes(blob['content']))
                state = 'indexed' if text is not None else 'not_indexed'
                con.execute('INSERT OR IGNORE INTO project_file_index(file_id,project_id,extracted_text,state,indexed_at) VALUES(?,?,?,?,?)',
                            (blob['id'], blob['project_id'], text or '', state, _now() if text is not None else None))
                con.execute('INSERT OR IGNORE INTO project_file_versions(file_id,version,media_type,content,size_bytes,created_at,created_by) VALUES(?,?,?,?,?,?,?)',
                            (blob['id'], 1, blob['media_type'], blob['content'], blob['size_bytes'], _now(), 'owner'))
                if text is not None:
                    try: con.execute('INSERT INTO project_file_fts(file_id,project_id,title,body) VALUES(?,?,?,?)', (blob['id'], blob['project_id'], blob['title'], text))
                    except sqlite3.OperationalError: pass

    @staticmethod
    def _dict(row):
        return dict(row) if row else None

    def _activity(self, con, project_id, action, target_type, target_id=None, detail='', actor='owner'):
        con.execute('INSERT INTO project_activity VALUES(?,?,?,?,?,?,?,?)',
                    (str(uuid.uuid4()), project_id, actor, action, target_type, target_id, detail[:1000], _now()))
        con.execute('UPDATE projects SET updated_at=? WHERE id=?', (_now(), project_id))

    def list(self, query='', status='all', sort='recent'):
        query = f'%{query.strip()}%'
        clauses = ['(p.name LIKE ? OR p.goal LIKE ? OR p.description LIKE ?)']
        values = [query, query, query]
        if status == 'all':
            clauses.append("p.status!='archived'")
        elif status != 'all':
            if status == 'needs_review':
                clauses.append("EXISTS (SELECT 1 FROM project_tasks r WHERE r.project_id=p.id AND r.status='needs_review')")
            elif status == 'completed':
                clauses.append("(p.status='completed' OR (EXISTS (SELECT 1 FROM project_tasks c WHERE c.project_id=p.id) AND NOT EXISTS (SELECT 1 FROM project_tasks c WHERE c.project_id=p.id AND c.status!='done')))")
            else:
                clauses.append('p.status=?')
                values.append(status)
        order = 'p.updated_at DESC' if sort == 'recent' else 'p.name COLLATE NOCASE ASC' if sort == 'name' else 'p.created_at DESC'
        with self.con() as con:
            rows = con.execute('''SELECT p.*,
                    (SELECT count(*) FROM project_tasks t WHERE t.project_id=p.id) task_count,
                    (SELECT count(*) FROM project_tasks t WHERE t.project_id=p.id AND t.status='done') done_count,
                    (SELECT count(*) FROM project_milestones m WHERE m.project_id=p.id) milestone_count,
                    (SELECT count(*) FROM project_tasks t WHERE t.project_id=p.id AND t.status='needs_review') review_count,
                    (SELECT title FROM project_tasks t WHERE t.project_id=p.id AND t.status IN ('in_progress','needs_review','blocked') ORDER BY t.updated_at DESC LIMIT 1) current_task
                FROM projects p WHERE '''+' AND '.join(clauses)+f' ORDER BY {order}', values).fetchall()
        return [dict(row) for row in rows]

    def tasks_for_day(self, day: str, *, limit: int = 500):
        with self.con() as con:
            rows = con.execute('''SELECT t.*,p.name AS project_name,p.status AS project_status
                FROM project_tasks t JOIN projects p ON p.id=t.project_id
                WHERE t.due_date=? AND p.status NOT IN ('archived','completed')
                ORDER BY t.due_date,t.created_at,t.id LIMIT ?''',
                (str(day), max(1, min(int(limit), 500)))).fetchall()
        return [dict(row) for row in rows]

    def activity_feed(self, *, limit: int = 200):
        with self.con() as con:
            rows = con.execute('''SELECT a.*,p.name AS project_name
                FROM project_activity a JOIN projects p ON p.id=a.project_id
                ORDER BY a.created_at DESC,a.id DESC LIMIT ?''',
                (max(1, min(int(limit), 500)),)).fetchall()
        return [dict(row) for row in rows]

    def create(self, *, name, goal='', description='', target_date=None, project_type='software', success_criteria='', instructions='', context_notes='', tasks=None, milestones=None):
        project_id, stamp = str(uuid.uuid4()), _now()
        with self.lock, self.con() as con:
            con.execute('''INSERT INTO projects(id,name,goal,description,target_date,project_type,success_criteria,instructions,context_notes,created_at,updated_at)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?)''',
                        (project_id, name.strip(), goal.strip(), description.strip(), target_date, project_type, success_criteria.strip(), instructions.strip(), context_notes.strip(), stamp, stamp))
            self._activity(con, project_id, 'created project', 'project', project_id, name.strip())
            milestone_ids = {}
            for item in milestones or []:
                milestone_id = str(uuid.uuid4())
                milestone_ids[item.get('client_id')] = milestone_id
                con.execute('''INSERT INTO project_milestones(id,project_id,title,outcome,target_date,completion_criteria,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?)''', (milestone_id, project_id, item['title'].strip(), item.get('outcome','').strip(), item.get('target_date'), item.get('completion_criteria','').strip(), stamp, stamp))
            for item in tasks or []:
                milestone_id = milestone_ids.get(item.get('milestone_client_id'))
                task_id = str(uuid.uuid4())
                con.execute('''INSERT INTO project_tasks(id,project_id,title,description,context_notes,status,priority,owner,due_date,milestone_id,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?)''', (task_id, project_id, item['title'].strip(), item.get('description','').strip(), item.get('context_notes','').strip(), 'planned', item.get('priority','medium'), item.get('owner','owner'), item.get('due_date'), milestone_id, stamp, stamp))
                if milestone_id:
                    con.execute('INSERT INTO project_milestone_tasks VALUES(?,?,?)', (milestone_id, task_id, project_id))
        return self.get(project_id)

    def create_walkthrough(self):
        """Create clearly labeled, persisted learning records for the signed-in owner."""
        with self.con() as con:
            existing = con.execute("SELECT id FROM projects WHERE is_walkthrough=1 AND status!='archived' ORDER BY created_at DESC LIMIT 1").fetchone()
        if existing:
            return self.get(existing['id'])
        target = (datetime.now(timezone.utc).date()).isoformat()
        project = self.create(name='Walkthrough · Personal AI redesign',
            goal='Use this sample workspace to learn how project goals, tasks, sources, and focused discussions fit together.',
            description='illustrative walkthrough records created on request. These examples do not represent actual work or Vishnu execution.',
            project_type='software', success_criteria='Understand the saved project workflow and the available project sections.')
        with self.lock, self.con() as con:
            con.execute('UPDATE projects SET is_walkthrough=1 WHERE id=?', (project['id'],))
        project = self.add_milestone(project['id'], title='Workspace review', outcome='Review project setup, context, and the saved work plan.',
                                     target_date=target, completion_criteria='Review each walkthrough section and understand what is persisted.')
        milestone_id = project['milestones'][0]['id']
        project = self.add_task(project['id'], title='Review the project overview', description='Inspect the goal, status summary, and project context.', owner='owner', milestone_id=milestone_id)
        overview_task_id = project['tasks'][0]['id']
        project = self.add_task(project['id'], title='Check the source indexing state', description='Compare a supported text source with an uploaded file that is stored only.', owner='owner', milestone_id=milestone_id)
        source_task_id = next(task['id'] for task in project['tasks'] if task['title']=='Check the source indexing state')
        project = self.update_task(project['id'], overview_task_id, {'status':'done'})
        project = self.update_task(project['id'], source_task_id, {'status':'needs_review'})
        project = self.add_upload(project['id'], title='walkthrough-brief.md', media_type='text/markdown',
            content=b'# Walkthrough brief\n\nThis is illustrative sample context created to demonstrate the Personal AI project workspace. It is not a real customer document.\n\n## Goal\nLearn how a goal, task, milestone, source file, and discussion connect.\n')
        file_id = next(file['id'] for file in project['files'] if file['title']=='walkthrough-brief.md')
        thread = self.add_thread(project['id'], title='Sample decision · keep project context scoped',
            content='Walkthrough example: project files and task context belong to this workspace.',
            linked_item_type='project_brief', linked_item_id=project['id'])['threads'][0]
        reply = thread['replies'][0]
        self.update_thread(project['id'], thread['id'], {'pinned_reply_id':reply['id'], 'followed':True})
        self.add_thread(project['id'], title='Review source support', content='Walkthrough example: Markdown can be supplied as project text context.',
            linked_item_type='file', linked_item_id=file_id)
        return self.get(project['id'])

    def get(self, project_id, *, include_archived=False):
        with self.con() as con:
            project = self._dict(con.execute('''SELECT p.*,
                    (SELECT count(*) FROM project_tasks t WHERE t.project_id=p.id) task_count,
                    (SELECT count(*) FROM project_tasks t WHERE t.project_id=p.id AND t.status='done') done_count,
                    (SELECT count(*) FROM project_milestones m WHERE m.project_id=p.id) milestone_count
                FROM projects p WHERE p.id=? AND (? OR p.status!='archived') ''', (project_id, int(include_archived))).fetchone())
            if not project:
                return None
            for key, table, order in [
                ('tasks', 'project_tasks', 'created_at'), ('milestones', 'project_milestones', 'target_date'),
                ('threads', 'project_threads', 'updated_at'), ('activity', 'project_activity', 'created_at DESC')]:
                project[key] = [dict(row) for row in con.execute(
                    f'SELECT * FROM {table} WHERE project_id=? ORDER BY {order}', (project_id,)).fetchall()]
            for task in project['tasks']:
                task['file_ids'] = [row['file_id'] for row in con.execute(
                    'SELECT file_id FROM project_task_files WHERE task_id=? ORDER BY file_id', (task['id'],)).fetchall()]
            project['files'] = [dict(row) for row in con.execute('''SELECT f.*,
                    (SELECT size_bytes FROM project_file_blobs b WHERE b.file_id=f.id) size_bytes,
                    COALESCE(i.state, CASE WHEN f.kind='link' THEN 'not_indexed' ELSE 'not_indexed' END) indexing_state,
                    (SELECT count(*) FROM project_file_versions v WHERE v.file_id=f.id) version_count,
                    (SELECT max(version) FROM project_file_versions v WHERE v.file_id=f.id) current_version
                FROM project_files f LEFT JOIN project_file_index i ON i.file_id=f.id
                WHERE f.project_id=? ORDER BY f.created_at''', (project_id,)).fetchall()]
            for thread in project['threads']:
                thread['replies'] = [dict(row) for row in con.execute(
                    'SELECT * FROM project_replies WHERE thread_id=? ORDER BY created_at', (thread['id'],)).fetchall()]
                thread['followed'] = bool(thread.get('followed'))
                thread['pinned_reply'] = next((reply for reply in thread['replies']
                                                if reply['id'] == thread.get('pinned_reply_id')), None)
                linked_type, linked_id = thread.get('linked_item_type'), thread.get('linked_item_id')
                if linked_type and linked_id:
                    source = {'project_brief': ('projects', 'id', 'name'), 'task': ('project_tasks', 'id', 'title'),
                              'file': ('project_files', 'id', 'title'), 'milestone': ('project_milestones', 'id', 'title'),
                              'proposal': ('project_approvals', 'id', 'title')}.get(linked_type)
                    if source:
                        table, key, label = source
                        if linked_type == 'project_brief':
                            linked = con.execute('SELECT name title FROM projects WHERE id=?', (project_id,)).fetchone()
                        else:
                            linked = con.execute(f'SELECT {label} title FROM {table} WHERE {key}=? AND project_id=?',
                                                  (linked_id, project_id)).fetchone()
                        thread['linked_item_title'] = linked['title'] if linked else ''
            for milestone in project['milestones']:
                milestone['task_ids'] = [row['task_id'] for row in con.execute(
                    'SELECT task_id FROM project_milestone_tasks WHERE milestone_id=? ORDER BY task_id', (milestone['id'],)).fetchall()]
            project['approvals'] = [dict(row) for row in con.execute(
                'SELECT * FROM project_approvals WHERE project_id=? ORDER BY proposed_at DESC', (project_id,)).fetchall()]
            for approval in project['approvals']:
                approval['task_ids'] = [row['task_id'] for row in con.execute('SELECT task_id FROM project_approval_tasks WHERE approval_id=?', (approval['id'],)).fetchall()]
                approval['file_ids'] = [row['file_id'] for row in con.execute('SELECT file_id FROM project_approval_files WHERE approval_id=?', (approval['id'],)).fetchall()]
                approval['history'] = [dict(row) for row in con.execute('SELECT actor,decision,comments,created_at FROM project_approval_history WHERE approval_id=? ORDER BY created_at', (approval['id'],)).fetchall()]
        return project

    def update(self, project_id, fields):
        allowed = {'name', 'goal', 'description', 'target_date', 'conversation_id', 'project_type', 'success_criteria', 'instructions', 'context_notes', 'status'}
        fields = {key: (value.strip() if isinstance(value, str) else value) for key, value in fields.items() if key in allowed}
        if not fields:
            return self.get(project_id)
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM projects WHERE id=?', (project_id,)).fetchone():
                return None
            con.execute('UPDATE projects SET '+','.join(f'{key}=?' for key in fields)+',updated_at=? WHERE id=?',
                        (*fields.values(), _now(), project_id))
            self._activity(con, project_id, 'updated project', 'project', project_id, ', '.join(fields))
        return self.get(project_id)

    def archive(self, project_id):
        with self.lock, self.con() as con:
            row = con.execute('SELECT id FROM projects WHERE id=? AND status!=\'archived\'', (project_id,)).fetchone()
            if not row:
                return False
            self._activity(con, project_id, 'archived project', 'project', project_id)
            con.execute('UPDATE projects SET status=\'archived\',updated_at=? WHERE id=?', (_now(), project_id))
        return True

    def restore(self, project_id):
        with self.lock, self.con() as con:
            row = con.execute("SELECT id FROM projects WHERE id=? AND status='archived'", (project_id,)).fetchone()
            if not row:
                return False
            con.execute("UPDATE projects SET status='active',updated_at=? WHERE id=?", (_now(), project_id))
            self._activity(con, project_id, 'restored project', 'project', project_id)
        return True

    def add_task(self, project_id, *, title, description='', context_notes='', status='planned', priority='medium', owner='vishnu', due_date=None, milestone_id=None, file_ids=None):
        task_id, stamp = str(uuid.uuid4()), _now()
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM projects WHERE id=? AND status!=\'archived\'', (project_id,)).fetchone():
                return None
            if milestone_id and not con.execute('SELECT 1 FROM project_milestones WHERE id=? AND project_id=?', (milestone_id, project_id)).fetchone():
                raise ValueError('Milestone does not belong to this project')
            self._validate_project_files(con, project_id, file_ids or [])
            con.execute('''INSERT INTO project_tasks(id,project_id,title,description,context_notes,status,priority,owner,due_date,milestone_id,created_at,updated_at)
                           VALUES(?,?,?,?,?,?,?,?,?,?,?,?)''', (task_id, project_id, title.strip(), description.strip(), context_notes.strip(), status, priority, owner, due_date, milestone_id, stamp, stamp))
            if milestone_id:
                con.execute('INSERT OR IGNORE INTO project_milestone_tasks VALUES(?,?,?)', (milestone_id, task_id, project_id))
            for file_id in file_ids or []:
                con.execute('INSERT INTO project_task_files VALUES(?,?,?)', (task_id, file_id, project_id))
            self._activity(con, project_id, 'created task', 'task', task_id, title.strip())
        return self.get(project_id)

    def update_task(self, project_id, task_id, fields):
        allowed = {'title', 'description', 'context_notes', 'status', 'priority', 'owner', 'due_date', 'milestone_id', 'file_ids'}
        fields = {key: (value.strip() if isinstance(value, str) else value) for key, value in fields.items() if key in allowed}
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM project_tasks WHERE id=? AND project_id=?', (task_id, project_id)).fetchone():
                return None
            if fields.get('milestone_id') and not con.execute('SELECT 1 FROM project_milestones WHERE id=? AND project_id=?', (fields['milestone_id'], project_id)).fetchone():
                raise ValueError('Milestone does not belong to this project')
            if fields:
                file_ids = fields.pop('file_ids', None)
                if file_ids is not None:
                    self._validate_project_files(con, project_id, file_ids)
                    con.execute('DELETE FROM project_task_files WHERE task_id=?', (task_id,))
                    for file_id in file_ids:
                        con.execute('INSERT INTO project_task_files VALUES(?,?,?)', (task_id, file_id, project_id))
                    con.execute('UPDATE project_tasks SET updated_at=? WHERE id=? AND project_id=?', (_now(), task_id, project_id))
                if fields:
                    con.execute('UPDATE project_tasks SET '+','.join(f'{key}=?' for key in fields)+',updated_at=? WHERE id=? AND project_id=?',
                                (*fields.values(), _now(), task_id, project_id))
                self._activity(con, project_id, 'updated task', 'task', task_id, ', '.join(fields))
                if 'milestone_id' in fields:
                    con.execute('DELETE FROM project_milestone_tasks WHERE task_id=? AND project_id=?', (task_id, project_id))
                    if fields['milestone_id']:
                        con.execute('INSERT OR IGNORE INTO project_milestone_tasks VALUES(?,?,?)', (fields['milestone_id'], task_id, project_id))
        return self.get(project_id)

    def set_task_execution(self, project_id, task_id, *, workflow_id=None, run_id=None, result=None, status=None):
        fields = {'execution_workflow_id': workflow_id, 'execution_run_id': run_id, 'execution_result': result}
        with self.lock, self.con() as con:
            row = con.execute('SELECT title FROM project_tasks WHERE id=? AND project_id=?', (task_id, project_id)).fetchone()
            if not row:
                return None
            updates = {key: value for key, value in fields.items() if value is not None}
            if status is not None:
                updates['status'] = status
            if updates:
                con.execute('UPDATE project_tasks SET '+','.join(f'{key}=?' for key in updates)+',updated_at=? WHERE id=? AND project_id=?',
                            (*updates.values(), _now(), task_id, project_id))
                action = 'started Vishnu task' if run_id else 'updated Vishnu task execution'
                self._activity(con, project_id, action, 'task', task_id, row['title'])
        return self.get(project_id)

    @staticmethod
    def _validate_project_files(con, project_id, file_ids):
        if not isinstance(file_ids, list) or len(file_ids) > 30 or len(file_ids) != len(set(file_ids)):
            raise ValueError('Choose up to 30 unique project files')
        for file_id in file_ids:
            if not con.execute('SELECT 1 FROM project_files WHERE id=? AND project_id=?', (file_id, project_id)).fetchone():
                raise ValueError('An attached file does not belong to this project')

    def add_approval(self, project_id, *, title, summary, impact_summary='', scope_summary='', task_ids=None, file_ids=None):
        approval_id, stamp = str(uuid.uuid4()), _now()
        with self.lock, self.con() as con:
            if not con.execute("SELECT 1 FROM projects WHERE id=? AND status!='archived'", (project_id,)).fetchone():
                return None
            self._validate_project_relations(con, project_id, task_ids or [], file_ids or [])
            con.execute('INSERT INTO project_approvals(id,project_id,title,summary,impact_summary,scope_summary,proposed_at,updated_at) VALUES(?,?,?,?,?,?,?,?)',
                        (approval_id, project_id, title.strip(), summary.strip(), impact_summary.strip(), scope_summary.strip(), stamp, stamp))
            for task_id in task_ids or []: con.execute('INSERT INTO project_approval_tasks VALUES(?,?)', (approval_id, task_id))
            for file_id in file_ids or []: con.execute('INSERT INTO project_approval_files VALUES(?,?)', (approval_id, file_id))
            self._activity(con, project_id, 'submitted proposal for review', 'proposal', approval_id, title.strip())
        return self.get(project_id)

    @staticmethod
    def _validate_project_relations(con, project_id, task_ids, file_ids):
        if len(task_ids) > 100 or len(task_ids) != len(set(task_ids)) or len(file_ids) > 100 or len(file_ids) != len(set(file_ids)):
            raise ValueError('Proposal links must be unique and within the supported limit')
        for task_id in task_ids:
            if not con.execute('SELECT 1 FROM project_tasks WHERE id=? AND project_id=?', (task_id, project_id)).fetchone():
                raise ValueError('A linked task does not belong to this project')
        ProjectStore._validate_project_files(con, project_id, file_ids)

    def revise_approval(self, project_id, approval_id, *, title, summary, impact_summary='', scope_summary='', task_ids=None, file_ids=None):
        stamp = _now()
        with self.lock, self.con() as con:
            row = con.execute("SELECT status FROM project_approvals WHERE id=? AND project_id=?", (approval_id, project_id)).fetchone()
            if not row: return None
            if row['status'] != 'changes_requested': raise ValueError('Only proposals with requested changes can be resubmitted')
            self._validate_project_relations(con, project_id, task_ids or [], file_ids or [])
            con.execute("UPDATE project_approvals SET title=?,summary=?,impact_summary=?,scope_summary=?,status='pending',reviewer=NULL,reviewed_at=NULL,decision=NULL,reviewer_comments='',updated_at=? WHERE id=?", (title.strip(), summary.strip(), impact_summary.strip(), scope_summary.strip(), stamp, approval_id))
            con.execute('DELETE FROM project_approval_tasks WHERE approval_id=?', (approval_id,)); con.execute('DELETE FROM project_approval_files WHERE approval_id=?', (approval_id,))
            for task_id in task_ids or []: con.execute('INSERT INTO project_approval_tasks VALUES(?,?)', (approval_id, task_id))
            for file_id in file_ids or []: con.execute('INSERT INTO project_approval_files VALUES(?,?)', (approval_id, file_id))
            con.execute('INSERT INTO project_approval_history VALUES(?,?,?,?,?,?)', (str(uuid.uuid4()), approval_id, 'owner', 'resubmitted', '', stamp))
            self._activity(con, project_id, 'resubmitted proposal', 'proposal', approval_id, title.strip())
        return self.get(project_id)

    def decide_approval(self, project_id, approval_id, *, decision, comments=''):
        if decision not in {'approved', 'changes_requested'}: raise ValueError('Choose approve or request changes')
        if decision == 'changes_requested' and not comments.strip(): raise ValueError('Add reviewer comments when requesting changes')
        stamp = _now()
        with self.lock, self.con() as con:
            row = con.execute("SELECT status,title FROM project_approvals WHERE id=? AND project_id=?", (approval_id, project_id)).fetchone()
            if not row: return None
            if row['status'] != 'pending': raise ValueError('This proposal has already been reviewed')
            con.execute('UPDATE project_approvals SET status=?,reviewer=?,reviewed_at=?,decision=?,reviewer_comments=?,updated_at=? WHERE id=?',
                        (decision, 'owner', stamp, decision, comments.strip(), stamp, approval_id))
            con.execute('INSERT INTO project_approval_history VALUES(?,?,?,?,?,?)', (str(uuid.uuid4()), approval_id, 'owner', decision, comments.strip(), stamp))
            self._activity(con, project_id, 'reviewed proposal', 'proposal', approval_id, decision)
        return self.get(project_id)

    def add_milestone(self, project_id, *, title, outcome='', target_date=None, completion_criteria='', task_ids=None):
        item_id, stamp = str(uuid.uuid4()), _now()
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM projects WHERE id=? AND status!=\'archived\'', (project_id,)).fetchone():
                return None
            con.execute('INSERT INTO project_milestones(id,project_id,title,outcome,target_date,completion_criteria,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)',
                        (item_id, project_id, title.strip(), outcome.strip(), target_date, completion_criteria.strip(), stamp, stamp))
            for task_id in task_ids or []:
                if not con.execute('SELECT 1 FROM project_tasks WHERE id=? AND project_id=?', (task_id, project_id)).fetchone():
                    raise ValueError('A linked task does not belong to this project')
                con.execute('DELETE FROM project_milestone_tasks WHERE task_id=? AND project_id=?', (task_id, project_id))
                con.execute('UPDATE project_tasks SET milestone_id=? WHERE id=? AND project_id=?', (item_id, task_id, project_id))
                con.execute('INSERT OR IGNORE INTO project_milestone_tasks VALUES(?,?,?)', (item_id, task_id, project_id))
            self._activity(con, project_id, 'created milestone', 'milestone', item_id, title.strip())
        return self.get(project_id)

    def update_milestone(self, project_id, milestone_id, fields):
        allowed = {'title', 'outcome', 'status', 'target_date', 'completion_criteria', 'task_ids'}
        fields = {key: (value.strip() if isinstance(value, str) else value) for key, value in fields.items() if key in allowed}
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM project_milestones WHERE id=? AND project_id=?', (milestone_id, project_id)).fetchone():
                return None
            if fields:
                task_ids = fields.pop('task_ids', None)
                if fields:
                    con.execute('UPDATE project_milestones SET '+','.join(f'{key}=?' for key in fields)+',updated_at=? WHERE id=? AND project_id=?',
                                (*fields.values(), _now(), milestone_id, project_id))
                if task_ids is not None:
                    for task_id in task_ids:
                        if not con.execute('SELECT 1 FROM project_tasks WHERE id=? AND project_id=?', (task_id, project_id)).fetchone():
                            raise ValueError('A linked task does not belong to this project')
                    con.execute('DELETE FROM project_milestone_tasks WHERE milestone_id=?', (milestone_id,))
                    con.execute('UPDATE project_tasks SET milestone_id=NULL WHERE milestone_id=? AND project_id=?', (milestone_id, project_id))
                    for task_id in task_ids:
                        con.execute('DELETE FROM project_milestone_tasks WHERE task_id=? AND project_id=?', (task_id, project_id))
                        con.execute('UPDATE project_tasks SET milestone_id=? WHERE id=? AND project_id=?', (milestone_id, task_id, project_id))
                        con.execute('INSERT INTO project_milestone_tasks VALUES(?,?,?)', (milestone_id, task_id, project_id))
                self._activity(con, project_id, 'updated milestone', 'milestone', milestone_id, ', '.join(fields))
        return self.get(project_id)

    def add_file(self, project_id, *, title, url, kind='link', knowledge_id=None):
        file_id, stamp = str(uuid.uuid4()), _now()
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM projects WHERE id=? AND status!=\'archived\'', (project_id,)).fetchone():
                return None
            con.execute('INSERT INTO project_files(id,project_id,title,kind,url,knowledge_id,created_at) VALUES(?,?,?,?,?,?,?)',
                        (file_id, project_id, title.strip(), kind, url.strip(), knowledge_id, stamp))
            self._activity(con, project_id, 'added project source', 'file', file_id, title.strip())
        return self.get(project_id)

    def add_upload(self, project_id, *, title, media_type, content):
        file_id, stamp = str(uuid.uuid4()), _now()
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM projects WHERE id=? AND status!=\'archived\'', (project_id,)).fetchone():
                return None
            con.execute('INSERT INTO project_files(id,project_id,title,kind,url,created_at) VALUES(?,?,?,?,?,?)',
                        (file_id, project_id, title.strip(), 'upload', '', stamp))
            con.execute('INSERT INTO project_file_blobs VALUES(?,?,?,?)', (file_id, media_type, sqlite3.Binary(content), len(content)))
            con.execute('INSERT INTO project_file_versions VALUES(?,?,?,?,?,?,?)', (file_id, 1, media_type, sqlite3.Binary(content), len(content), stamp, 'owner'))
            text = self._extract_text(title, media_type, content)
            state = 'indexed' if text is not None else 'not_indexed'
            con.execute('INSERT INTO project_file_index(file_id,project_id,extracted_text,state,indexed_at) VALUES(?,?,?,?,?)',
                        (file_id, project_id, text or '', state, stamp if text is not None else None))
            if text is not None:
                try:
                    con.execute('INSERT INTO project_file_fts(file_id,project_id,title,body) VALUES(?,?,?,?)', (file_id, project_id, title.strip(), text))
                except sqlite3.OperationalError:
                    pass
            self._activity(con, project_id, 'uploaded project file', 'file', file_id, title.strip())
        return self.get(project_id)

    def add_drive_source(self, project_id, *, title, media_type, content, provider_file_id, modified_at=''):
        if len(content) > 10 * 1024 * 1024 or not content:
            raise ValueError('Google Drive files must be between 1 byte and 10 MB')
        stamp, file_id = _now(), str(uuid.uuid4())
        with self.lock, self.con() as con:
            if not con.execute("SELECT 1 FROM projects WHERE id=? AND status!='archived'", (project_id,)).fetchone(): return None
            existing = con.execute("SELECT id FROM project_files WHERE project_id=? AND source_provider='google_drive' AND source_file_id=?",
                                   (project_id, str(provider_file_id))).fetchone()
            if existing:
                # A second explicit import refreshes the snapshot and records a new immutable version.
                file_id = existing['id']
        if existing:
            return self.update_drive_source(project_id, file_id, title=title, media_type=media_type, content=content, modified_at=modified_at)
        with self.lock, self.con() as con:
            con.execute("INSERT INTO project_files(id,project_id,title,kind,url,created_at,source_provider,source_file_id,source_modified_at) VALUES(?,?,?,?,?,?,?,?,?)",
                        (file_id, project_id, str(title)[:180], 'upload', '', stamp, 'google_drive', str(provider_file_id), str(modified_at or '')))
            con.execute('INSERT INTO project_file_blobs VALUES(?,?,?,?)', (file_id, media_type, sqlite3.Binary(content), len(content)))
            con.execute('INSERT INTO project_file_versions VALUES(?,?,?,?,?,?,?)', (file_id, 1, media_type, sqlite3.Binary(content), len(content), stamp, 'owner'))
            text = self._extract_text(title, media_type, content)
            con.execute('INSERT INTO project_file_index(file_id,project_id,extracted_text,state,indexed_at) VALUES(?,?,?,?,?)',
                        (file_id, project_id, text or '', 'indexed' if text is not None else 'not_indexed', stamp if text is not None else None))
            if text is not None:
                try: con.execute('INSERT INTO project_file_fts(file_id,project_id,title,body) VALUES(?,?,?,?)', (file_id, project_id, title, text))
                except sqlite3.OperationalError: pass
            self._activity(con, project_id, 'imported Google Drive source', 'file', file_id, title[:180])
        return self.get(project_id)

    def update_drive_source(self, project_id, file_id, *, title, media_type, content, modified_at=''):
        if len(content) > 10 * 1024 * 1024 or not content:
            raise ValueError('Google Drive files must be between 1 byte and 10 MB')
        stamp = _now()
        with self.lock, self.con() as con:
            row = con.execute("SELECT source_file_id FROM project_files WHERE id=? AND project_id=? AND source_provider='google_drive'",
                              (file_id, project_id)).fetchone()
            if not row:
                return None
            next_version = int(con.execute('SELECT COALESCE(max(version),0)+1 FROM project_file_versions WHERE file_id=?', (file_id,)).fetchone()[0])
            con.execute('UPDATE project_files SET title=?,source_modified_at=?,created_at=? WHERE id=? AND project_id=?',
                        (str(title)[:180], str(modified_at or ''), stamp, file_id, project_id))
            con.execute('UPDATE project_file_blobs SET media_type=?,content=?,size_bytes=? WHERE file_id=?',
                        (media_type, sqlite3.Binary(content), len(content), file_id))
            con.execute('INSERT INTO project_file_versions VALUES(?,?,?,?,?,?,?)',
                        (file_id, next_version, media_type, sqlite3.Binary(content), len(content), stamp, 'owner'))
            text = self._extract_text(title, media_type, content)
            con.execute('UPDATE project_file_index SET extracted_text=?,state=?,indexed_at=? WHERE file_id=?',
                        (text or '', 'indexed' if text is not None else 'not_indexed', stamp if text is not None else None, file_id))
            try:
                con.execute('DELETE FROM project_file_fts WHERE file_id=?', (file_id,))
                if text is not None: con.execute('INSERT INTO project_file_fts(file_id,project_id,title,body) VALUES(?,?,?,?)', (file_id, project_id, str(title)[:180], text))
            except sqlite3.OperationalError: pass
            self._activity(con, project_id, 'refreshed Google Drive source', 'file', file_id, f'{title} · snapshot {next_version}')
        return self.get(project_id)

    @staticmethod
    def _extract_text(title, media_type, content):
        suffix = str(title).rsplit('.', 1)[-1].lower() if '.' in str(title) else ''
        if len(content) > 12 * 1024 * 1024:
            return None
        try:
            if suffix in {'txt', 'md', 'csv'}:
                return content.decode('utf-8-sig')[:2_000_000]
            if suffix == 'pdf':
                from pypdf import PdfReader
                import io
                reader = PdfReader(io.BytesIO(content), strict=False)
                if len(reader.pages) > 500: return None
                return '\n'.join((page.extract_text() or '') for page in reader.pages)[:2_000_000]
            if suffix == 'docx':
                from docx import Document
                import io
                doc = Document(io.BytesIO(content))
                return '\n'.join(p.text for p in doc.paragraphs)[:2_000_000]
            if suffix == 'xlsx':
                from openpyxl import load_workbook
                import io
                book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                values=[]
                for sheet in book.worksheets[:20]:
                    values.append(f'[{sheet.title}]')
                    for row_num,row in enumerate(sheet.iter_rows(values_only=True), start=1):
                        if row_num > 5000: break
                        values.append('\t'.join('' if value is None else str(value)[:500] for value in row[:50]))
                        if sum(map(len, values)) > 2_000_000: break
                    if sum(map(len, values)) > 2_000_000: break
                book.close()
                return '\n'.join(values)[:2_000_000]
        except Exception:
            return None
        return None

    def search_files(self, project_id, query, *, limit=50):
        query = str(query or '').strip()[:200]
        if not query:
            return []
        limit = max(1, min(int(limit), 100))
        terms = [part for part in query.split() if part]
        with self.con() as con:
            try:
                match = ' AND '.join('"' + term.replace('"', '""') + '"*' for term in terms)
                rows = con.execute('''SELECT f.id,f.title,f.kind,f.url,f.created_at,i.state indexing_state,
                    b.size_bytes FROM project_file_fts x JOIN project_files f ON f.id=x.file_id AND f.project_id=x.project_id
                    LEFT JOIN project_file_index i ON i.file_id=f.id LEFT JOIN project_file_blobs b ON b.file_id=f.id
                    WHERE x.project_id=? AND project_file_fts MATCH ? ORDER BY rank LIMIT ?''', (project_id, match, limit)).fetchall()
            except sqlite3.OperationalError:
                like = '%' + query.replace('%', '\\%').replace('_', '\\_') + '%'
                rows = con.execute('''SELECT f.id,f.title,f.kind,f.url,f.created_at,i.state indexing_state,b.size_bytes
                    FROM project_files f JOIN project_file_index i ON i.file_id=f.id
                    LEFT JOIN project_file_blobs b ON b.file_id=f.id WHERE f.project_id=? AND
                    (f.title LIKE ? ESCAPE '\\' OR i.extracted_text LIKE ? ESCAPE '\\') ORDER BY f.created_at DESC LIMIT ?''',
                    (project_id, like, like, limit)).fetchall()
        return [dict(row) for row in rows]

    def save_deliverable(self, project_id, *, title, content, file_id=None):
        title = str(title or '').strip()[:180]
        if not title or not isinstance(content, (str, bytes)):
            raise ValueError('A title and deliverable content are required')
        raw = content.encode('utf-8') if isinstance(content, str) else content
        if not raw or len(raw) > 1_000_000:
            raise ValueError('Deliverables must be between 1 byte and 1 MB')
        media_type = 'text/markdown; charset=utf-8'
        stamp = _now()
        with self.lock, self.con() as con:
            if not con.execute("SELECT 1 FROM projects WHERE id=? AND status!='archived'", (project_id,)).fetchone():
                return None
            if file_id:
                row = con.execute("SELECT id FROM project_files WHERE id=? AND project_id=? AND kind='deliverable'", (file_id, project_id)).fetchone()
                if not row:
                    return None
                next_version = int(con.execute('SELECT COALESCE(max(version),0)+1 FROM project_file_versions WHERE file_id=?', (file_id,)).fetchone()[0])
                try: con.execute('DELETE FROM project_file_fts WHERE file_id=?', (file_id,))
                except sqlite3.OperationalError: pass
                con.execute('UPDATE project_files SET title=?,created_at=?,deliverable_status=? WHERE id=? AND project_id=?', (title, stamp, 'ready_for_review', file_id, project_id))
                con.execute('UPDATE project_file_blobs SET media_type=?,content=?,size_bytes=? WHERE file_id=?', (media_type, sqlite3.Binary(raw), len(raw), file_id))
                con.execute('INSERT INTO project_file_versions VALUES(?,?,?,?,?,?,?)', (file_id, next_version, media_type, sqlite3.Binary(raw), len(raw), stamp, 'owner'))
                con.execute('DELETE FROM project_file_index WHERE file_id=?', (file_id,))
            else:
                file_id, next_version = str(uuid.uuid4()), 1
                con.execute("INSERT INTO project_files(id,project_id,title,kind,url,created_at,source_state,deliverable_status) VALUES(?,?,?,?,?,?,?,?)",
                            (file_id, project_id, title, 'deliverable', '', stamp, 'stored', 'ready_for_review'))
                con.execute('INSERT INTO project_file_blobs VALUES(?,?,?,?)', (file_id, media_type, sqlite3.Binary(raw), len(raw)))
                con.execute('INSERT INTO project_file_versions VALUES(?,?,?,?,?,?,?)', (file_id, next_version, media_type, sqlite3.Binary(raw), len(raw), stamp, 'owner'))
            text = raw.decode('utf-8')
            con.execute('INSERT INTO project_file_index(file_id,project_id,extracted_text,state,indexed_at) VALUES(?,?,?,?,?)', (file_id, project_id, text, 'indexed', stamp))
            try:
                con.execute('INSERT INTO project_file_fts(file_id,project_id,title,body) VALUES(?,?,?,?)', (file_id, project_id, title, text))
            except sqlite3.OperationalError:
                pass
            self._activity(con, project_id, 'saved project deliverable', 'file', file_id, f'{title} · version {next_version}')
        return self.get(project_id)

    def file_versions(self, project_id, file_id):
        with self.con() as con:
            valid = con.execute('SELECT 1 FROM project_files WHERE id=? AND project_id=?', (file_id, project_id)).fetchone()
            if not valid:
                return None
            return [dict(row) for row in con.execute('SELECT version,media_type,size_bytes,created_at,created_by FROM project_file_versions WHERE file_id=? ORDER BY version DESC', (file_id,)).fetchall()]

    def file_version_blob(self, project_id, file_id, version):
        with self.con() as con:
            row = con.execute('''SELECT f.title,v.version,v.media_type,v.content,v.size_bytes
                FROM project_files f JOIN project_file_versions v ON v.file_id=f.id
                WHERE f.id=? AND f.project_id=? AND v.version=?''',
                (file_id, project_id, version)).fetchone()
        return dict(row) if row else None

    def file_blob(self, project_id, file_id):
        with self.con() as con:
            row = con.execute('''SELECT f.title,b.media_type,b.content,b.size_bytes FROM project_files f
                JOIN project_file_blobs b ON b.file_id=f.id WHERE f.id=? AND f.project_id=?''', (file_id, project_id)).fetchone()
        return dict(row) if row else None

    def file_index(self, project_id, file_id):
        with self.con() as con:
            row = con.execute('''SELECT f.id,f.title,COALESCE(i.state,'not_indexed') state,
                COALESCE(i.extracted_text,'') extracted_text,COALESCE(i.error,'') error,i.indexed_at
                FROM project_files f LEFT JOIN project_file_index i ON i.file_id=f.id
                WHERE f.id=? AND f.project_id=?''', (file_id, project_id)).fetchone()
        return dict(row) if row else None

    def update_file(self, project_id, file_id, fields):
        if 'is_pinned' not in fields:
            return self.get(project_id)
        pinned = int(bool(fields['is_pinned']))
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM project_files WHERE id=? AND project_id=?', (file_id, project_id)).fetchone():
                return None
            con.execute('UPDATE project_files SET is_pinned=? WHERE id=? AND project_id=?', (pinned, file_id, project_id))
            self._activity(con, project_id, 'pinned source' if pinned else 'unpinned source', 'file', file_id)
        return self.get(project_id)

    @staticmethod
    def _validate_discussion_link(con, project_id, linked_item_type, linked_item_id):
        if not linked_item_type and not linked_item_id:
            return
        tables = {'project_brief': ('projects', 'id'), 'task': ('project_tasks', 'id'),
                  'file': ('project_files', 'id'), 'milestone': ('project_milestones', 'id'),
                  'proposal': ('project_approvals', 'id')}
        if linked_item_type not in tables or not linked_item_id:
            raise ValueError('Choose a valid linked item from this project')
        table, key = tables[linked_item_type]
        if linked_item_type == 'project_brief':
            valid = linked_item_id == project_id
        else:
            valid = con.execute(f'SELECT 1 FROM {table} WHERE {key}=? AND project_id=?',
                                (linked_item_id, project_id)).fetchone()
        if not valid:
            raise ValueError('Linked item does not belong to this project')

    def add_thread(self, project_id, *, title, content, linked_item_type='', linked_item_id=None):
        thread_id, stamp = str(uuid.uuid4()), _now()
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM projects WHERE id=? AND status!=\'archived\'', (project_id,)).fetchone():
                return None
            self._validate_discussion_link(con, project_id, linked_item_type, linked_item_id)
            con.execute('''INSERT INTO project_threads(id,project_id,title,status,created_at,updated_at,
                linked_item_type,linked_item_id) VALUES(?,?,?,?,?,?,?,?)''',
                (thread_id, project_id, title.strip(), 'open', stamp, stamp, linked_item_type or '', linked_item_id))
            con.execute('INSERT INTO project_replies VALUES(?,?,?,?,?)', (str(uuid.uuid4()), thread_id, 'owner', content.strip(), stamp))
            self._activity(con, project_id, 'started discussion', 'discussion', thread_id, title.strip())
        return self.get(project_id)

    def update_thread(self, project_id, thread_id, fields):
        allowed = {'title', 'status', 'linked_item_type', 'linked_item_id', 'pinned_reply_id', 'followed'}
        fields = {key: (value.strip() if isinstance(value, str) else value) for key, value in fields.items() if key in allowed}
        with self.lock, self.con() as con:
            row = con.execute('SELECT * FROM project_threads WHERE id=? AND project_id=?', (thread_id, project_id)).fetchone()
            if not row:
                return None
            if 'linked_item_type' in fields or 'linked_item_id' in fields:
                self._validate_discussion_link(con, project_id, fields.get('linked_item_type', row['linked_item_type']),
                                               fields.get('linked_item_id', row['linked_item_id']))
            if 'pinned_reply_id' in fields and fields['pinned_reply_id'] and not con.execute(
                'SELECT 1 FROM project_replies WHERE id=? AND thread_id=?', (fields['pinned_reply_id'], thread_id)).fetchone():
                raise ValueError('Pinned decision reply must belong to this discussion')
            if 'status' in fields and fields['status'] not in {'open', 'resolved'}:
                raise ValueError('Discussion status must be open or resolved')
            if 'followed' in fields:
                fields['followed'] = int(bool(fields['followed']))
            if 'status' in fields:
                if fields['status'] == 'resolved':
                    con.execute('UPDATE project_threads SET resolved_by=COALESCE(resolved_by,\'owner\'), resolved_at=COALESCE(resolved_at,?) WHERE id=?', (_now(), thread_id))
                else:
                    con.execute('UPDATE project_threads SET resolved_by=NULL,resolved_at=NULL WHERE id=?', (thread_id,))
            if fields:
                con.execute('UPDATE project_threads SET '+','.join(f'{key}=?' for key in fields)+',updated_at=? WHERE id=? AND project_id=?',
                            (*fields.values(), _now(), thread_id, project_id))
                self._activity(con, project_id, 'updated discussion', 'discussion', thread_id,
                               ', '.join(fields.keys()))
        return self.get(project_id)

    def add_reply(self, project_id, thread_id, *, author, content):
        stamp = _now()
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM project_threads WHERE id=? AND project_id=?', (thread_id, project_id)).fetchone():
                return None
            con.execute('INSERT INTO project_replies VALUES(?,?,?,?,?)', (str(uuid.uuid4()), thread_id, author, content.strip(), stamp))
            con.execute('UPDATE project_threads SET updated_at=? WHERE id=?', (stamp, thread_id))
            self._activity(con, project_id, 'replied to discussion', 'discussion', thread_id, content.strip()[:160], actor=author)
        return self.get(project_id)

    def record(self, project_id, *, actor, action, target_type='project', target_id=None, detail=''):
        with self.lock, self.con() as con:
            if not con.execute('SELECT 1 FROM projects WHERE id=? AND status!=\'archived\'', (project_id,)).fetchone():
                return False
            self._activity(con, project_id, action, target_type, target_id, detail, actor=actor)
        return True
